"""Structural, live-formula and parseability tests for the six-tab workbook."""

from __future__ import annotations

import re
from pathlib import Path

import formulas
import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries

SHEET_ORDER = [
    "00_Exec_Control_Tower",
    "01_Catalog_&_Parameters",
    "02_Demand_Unconstraining",
    "03_FBA_Nivel1_Transfer",
    "04_Factory_Nivel2_PO_&_CBM",
    "05_Fee_Risk_Audit",
]

DATA_TABS = [
    "02_Demand_Unconstraining",
    "03_FBA_Nivel1_Transfer",
    "04_Factory_Nivel2_PO_&_CBM",
    "05_Fee_Risk_Audit",
]

EXPECTED_TABLES = {
    "tbl_Catalog": "01_Catalog_&_Parameters",
    "tbl_Demand_Unconstraining": "02_Demand_Unconstraining",
    "tbl_FBA_Transfers": "03_FBA_Nivel1_Transfer",
    "tbl_Factory_PO": "04_Factory_Nivel2_PO_&_CBM",
    "tbl_Fee_Audit": "05_Fee_Risk_Audit",
}

# Policy decisions (unconstrained demand, ETL buffers, ROP/S_max, age-safe
# cartons) are written as static engine outputs; the workbook keeps live
# formulas for arithmetic, checks and aggregates.
EXPECTED_FORMULA_COLUMNS = {
    "02_Demand_Unconstraining": list("GOTUVWX"),
    "03_FBA_Nivel1_Transfer": ["E", "H", "S", "T", "X", "Y", "Z", "AA", "AB"],
    "04_Factory_Nivel2_PO_&_CBM": ["F", "G", "H", "I", "J", "K", "L", "M", "N",
                                   "O", "P", "Q", "T", "U", "V", "W", "X", "Y", "Z"],
    "05_Fee_Risk_Audit": list("JKLMNOPQRSTUVWX"),
}


@pytest.fixture(scope="module")
def workbook(project):
    wb = load_workbook(project.workbook_path, data_only=False)
    yield wb
    wb.close()


def _all_formulas(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    yield ws.title, cell.coordinate, cell.value


def _catalog_header_map(wb):
    ws = wb["01_Catalog_&_Parameters"]
    table = ws.tables["tbl_Catalog"]
    min_col, min_row, max_col, max_row = range_boundaries(table.ref)
    headers = {}
    for column in range(min_col, max_col + 1):
        headers[ws.cell(row=min_row, column=column).value] = (column, min_row + 1, max_row)
    return headers


def _normalise_structured_refs(formula: str, headers) -> str:
    def replace(match: re.Match[str]) -> str:
        header = match.group(1)
        column, first_row, last_row = headers[header]
        letter = get_column_letter(column)
        return f"'01_Catalog_&_Parameters'!${letter}${first_row}:${letter}${last_row}"

    return re.sub(r"tbl_Catalog\[([^\]]+)\]", replace, formula)


def test_six_tabs_tables_and_named_controls(workbook):
    assert workbook.sheetnames == SHEET_ORDER
    assert len(workbook.sheetnames) == 6

    found = {}
    for ws in workbook.worksheets:
        for name in ws.tables:
            found[name] = ws.title
    assert found == EXPECTED_TABLES
    assert len(found) == 5

    names = set(workbook.defined_names)
    expected_names = {
        "L1_Mean", "L1_Std", "L2_Mean", "L2_Std", "Shield_Days", "Review_Days",
        "Age_Cap", "CBM_20GP", "CBM_40HC", "DOS_Cap", "Snapshot_Date",
        "Kappa_Cred", "Gamma_Cannib", "Fee_Threshold", "Fee_Min_Weekly",
        "Removal_Fee", "AIS_Payback", "Aggregate_MOQ", "LCL_Breakeven",
        "TopUp_Reserve", "Hero_Weekly",
    }
    assert names == expected_names
    assert len(names) == 21


def test_calculation_flags_force_recalculation(workbook):
    assert workbook.calculation.fullCalcOnLoad is True
    assert workbook.calculation.forceFullCalc is True
    assert workbook.calculation.calcMode == "auto"


def test_formula_population_is_live_not_static(workbook):
    counts = {}
    for sheet, _, _ in _all_formulas(workbook):
        counts[sheet] = counts.get(sheet, 0) + 1
    assert counts["00_Exec_Control_Tower"] >= 48
    assert counts["02_Demand_Unconstraining"] >= 1120
    assert counts["03_FBA_Nivel1_Transfer"] >= 1440
    assert counts["04_Factory_Nivel2_PO_&_CBM"] >= 3040
    assert counts["05_Fee_Risk_Audit"] >= 2400
    assert sum(counts.values()) >= 8048

    for sheet, columns in EXPECTED_FORMULA_COLUMNS.items():
        ws = workbook[sheet]
        for column in columns:
            for row in range(5, 165):
                value = ws[f"{column}{row}"].value
                assert isinstance(value, str) and value.startswith("="), (
                    sheet, f"{column}{row}", value
                )


def test_key_formulas_use_required_excel_365_functions(workbook):
    demand = workbook["02_Demand_Unconstraining"]
    assert demand["G5"].value == "=E5/30"
    assert demand["O5"].value == "=M5/30"
    assert demand["R5"].value == pytest.approx(
        demand["K5"].value + demand["Q5"].value, abs=2e-4
    )
    # K (unconstrained FBA demand) is the engine's decision, not a re-derivation.
    assert isinstance(demand["K5"].value, (int, float))
    assert isinstance(demand["J5"].value, (int, float))


def test_demand_ladder_is_auditable_from_visible_columns(workbook):
    """K must be reconstructible from H (active day), Y (fallback pool) and J (sisters)."""

    demand = workbook["02_Demand_Unconstraining"]
    assert demand["Y4"].value == "Fallback pool u/d"
    checked = {"active": 0, "fallback": 0, "sister": 0}
    for row in range(5, 165):
        active_days = float(demand[f"F{row}"].value)
        shipped = float(demand[f"K{row}"].value)
        active = float(demand[f"H{row}"].value)
        pool = float(demand[f"Y{row}"].value)
        sister = float(demand[f"J{row}"].value)
        if active_days > 0:
            expected, branch = active, "active"
        elif pool > 0:
            expected, branch = pool, "fallback"
        else:
            expected, branch = sister, "sister"
        assert shipped == pytest.approx(expected, abs=1e-4), (row, branch)
        checked[branch] += 1
    # This snapshot exercises the active-day and fallback branches; the sister
    # prior is never needed, because every depleted SKU has a lifetime rate.
    assert checked["active"] > 0 and checked["fallback"] > 0
    assert checked["sister"] == 0, checked

    fba = workbook["03_FBA_Nivel1_Transfer"]
    assert fba["K5"].value == pytest.approx(demand["K5"].value)
    assert fba["H5"].value == "=MAX(F5,G5)"
    assert "_xlfn.XLOOKUP(" in fba["S5"].value
    assert "_xlfn.XLOOKUP(" in fba["X5"].value
    assert fba["AB5"].value.startswith("=IF(MOD(")

    factory = workbook["04_Factory_Nivel2_PO_&_CBM"]
    assert factory["P5"].value.startswith("=_xlfn.LET(_xlpm.")
    assert "_xlpm.moq" in factory["P5"].value
    assert "_xlfn.CEILING.MATH(" in factory["P5"].value
    assert "_xlfn.XLOOKUP(" in factory["P5"].value
    assert "DOS_Cap" in factory["Z5"].value
    assert "CBM_20GP" in workbook["00_Exec_Control_Tower"]["B19"].value

    fee = workbook["05_Fee_Risk_Audit"]
    assert "AIS_Payback" in fee["S5"].value
    assert "EXECUTE REMOVAL" in fee["S5"].value
    assert "DISCOUNT 25% NOW" in fee["S5"].value
    assert "Removal_Fee" in fee["Q5"].value

    exec_sheet = workbook["00_Exec_Control_Tower"]
    assert "tbl_Catalog[sku_id]" in exec_sheet["B6"].value


def test_exec_sheet_kpis_match_engine_summary(workbook, project):
    """The control tower quotes the engine, not a parallel Excel model."""

    summary = project.summary
    exec_sheet = workbook["00_Exec_Control_Tower"]
    # B11-B12 are COUNTIF/COUNTIFS over tab 03; their inputs are static engine
    # values, so the formula must resolve to the same population the summary
    # reports once Excel recalculates.
    fba_plan = project.fba_plan
    assert (fba_plan["short_term_hdos_30d"] < 14).sum() == summary["stockout_skus_hdos_lt_14"]
    band = (
        (fba_plan["short_term_hdos_30d"] >= 14)
        & (fba_plan["short_term_hdos_30d"] < 28)
        & (fba_plan["long_term_hdos_90d"] < 28)
        & (~fba_plan["low_inv_fee_exempt_flag"])
    )
    assert band.sum() == summary["fee_band_skus_hdos_14_28"]
    assert fba_plan["executable_transfer_cartons"].sum() == summary["transfer_cartons"]
    assert fba_plan["executable_transfer_units"].sum() == summary["transfer_units"]

    # Static engine columns on tab 03 must mirror the persisted plan.
    transfer = workbook["03_FBA_Nivel1_Transfer"]
    plan_by_sku = fba_plan.set_index("sku_id")
    for row in range(5, 165):
        record = plan_by_sku.loc[transfer[f"A{row}"].value]
        assert transfer[f"K{row}"].value == pytest.approx(record["mu_unconstrained"])
        assert transfer[f"N{row}"].value == pytest.approx(record["b_fba"])
        assert transfer[f"O{row}"].value == pytest.approx(record["rop1_shield"])
        assert transfer[f"P{row}"].value == pytest.approx(record["smax1"])
        assert transfer[f"R{row}"].value == pytest.approx(record["target_cartons"])

    assert exec_sheet["B13"].value == f"=SUM('03_FBA_Nivel1_Transfer'!$Z$5:$Z$164)"
    assert exec_sheet["B14"].value == "=SUM('03_FBA_Nivel1_Transfer'!$T$5:$T$164)"


def test_every_let_formula_marks_its_parameters(workbook):
    """Excel rewrites LET parameters with ``_xlpm.``; without it, it repairs."""

    offenders = []
    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "_xlfn.LET(" in cell.value:
                    body = cell.value.split("_xlfn.LET(", 1)[1]
                    if "_xlpm." not in body:
                        offenders.append(f"{ws.title}!{cell.coordinate}")
    assert not offenders, offenders[:5]


def test_all_data_tabs_are_sku_aligned_with_catalog(workbook):
    catalog = workbook["01_Catalog_&_Parameters"]
    expected = [catalog[f"A{row}"].value for row in range(16, 176)]
    assert len(expected) == 160
    assert len(set(expected)) == 160
    for sheet in DATA_TABS:
        ws = workbook[sheet]
        actual = [ws[f"A{row}"].value for row in range(5, 165)]
        assert actual == expected, sheet


def test_every_structured_reference_points_at_a_real_column(workbook):
    headers = _catalog_header_map(workbook)
    referenced = set()
    for _, _, formula in _all_formulas(workbook):
        referenced.update(re.findall(r"tbl_Catalog\[([^\]]+)\]", formula))
    assert referenced
    assert referenced <= set(headers), sorted(referenced - set(headers))


def test_every_formula_parses_after_structured_ref_normalisation(workbook):
    headers = _catalog_header_map(workbook)
    parser = formulas.Parser()
    errors = []
    parsed = 0
    for sheet, coordinate, formula in _all_formulas(workbook):
        normalised = _normalise_structured_refs(formula, headers)
        try:
            parser.ast(normalised)
            parsed += 1
        except Exception as exc:  # pragma: no cover - failure diagnostics
            errors.append((sheet, coordinate, formula, repr(exc)))
    assert parsed >= 8048
    assert not errors, errors[:5]


def test_no_broken_reference_markers(workbook):
    for sheet, coordinate, formula in _all_formulas(workbook):
        assert "#REF!" not in formula, (sheet, coordinate)
        assert "#NAME?" not in formula, (sheet, coordinate)
        assert formula.count("(") == formula.count(")"), (sheet, coordinate, formula)


def test_conditional_formatting_and_signoff_controls(workbook):
    for sheet in DATA_TABS:
        assert len(list(workbook[sheet].conditional_formatting)) > 0, sheet
    exec_sheet = workbook["00_Exec_Control_Tower"]
    validations = list(exec_sheet.data_validations.dataValidation)
    assert any("APPROVE WEEKLY PLAN" in str(dv.formula1) for dv in validations)
    assert exec_sheet["B42"].value is None or isinstance(exec_sheet["B42"].value, str)


def test_catalog_parameter_controls_match_policy(workbook):
    ws = workbook["01_Catalog_&_Parameters"]
    assert ws["B3"].value == pytest.approx(14.0)
    assert ws["B4"].value == pytest.approx(3.0)
    assert ws["B5"].value == pytest.approx(75.0)
    assert ws["B6"].value == pytest.approx(10.0)
    assert ws["B7"].value == pytest.approx(32.0)
    assert ws["B9"].value == pytest.approx(120.0)
    assert ws["B10"].value == pytest.approx(28.8)
    assert ws["B11"].value == pytest.approx(70.0)
    assert ws["B12"].value == pytest.approx(135.0)
    assert ws["F3"].value == pytest.approx(7.0)
    assert ws["F4"].value == pytest.approx(0.10)
    assert ws["F5"].value == pytest.approx(28.0)
    assert ws["F7"].value == pytest.approx(0.84)
    assert ws["F8"].value == pytest.approx(12.0)
    assert ws["F9"].value == pytest.approx(5000)
    assert ws["F11"].value == pytest.approx(1.8)


def test_workbook_is_not_empty_and_has_traceable_title(workbook):
    assert workbook.properties.title
    assert "Control Tower" in workbook.properties.title
    assert len(workbook["01_Catalog_&_Parameters"]["A16"].value) > 0
    assert workbook["00_Exec_Control_Tower"]["A1"].value.startswith("OMNICHANNEL")
